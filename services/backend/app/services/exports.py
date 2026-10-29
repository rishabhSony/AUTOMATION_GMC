"""
Reviewed Quote Excel (Plan Phase 8 / PRD §6.5): the 5-column quote form, the
hospitalization & treatment sub-tables, the 19 procedure sub-limits, and underwriter notes.
The workbook only formats data already reviewed in the workbench; it computes nothing.
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

QUOTE_FILENAME = "gmc_quote_summary_sheet.xlsx"
QUOTE_COLUMNS = ["Quote Fields", "Coverage Details", "Proposed Details", "Source Status",
                 "Confidence"]  # fmt: skip

SECTION_LABELS = {
    "quotation": "Quotation Details",
    "policyholder": "Policyholder Details",
    "group_details": "Group Details",
    "hospital_expenses": "Hospital Expenses",
    "medical_expenses": "Medical Expenses",
    "maternity_newborn": "Maternity & Newborn",
    "medical_coverage": "Medical Coverages",
    "modern_treatments": "Modern Treatments",
    "wellness": "Wellness",
    "policy_conditions": "Policy Conditions",
}
# Hospitalization & treatment sub-tables (App Flow /cases/[id]/treatments) get their own sheet.
TREATMENT_SECTIONS = {
    "hospital_expenses", "medical_expenses", "maternity_newborn", "medical_coverage",
    "modern_treatments", "wellness", "policy_conditions",
}  # fmt: skip

_HEADER_FILL = PatternFill("solid", fgColor="0052CC")  # Pacific Cobalt (UI brief)
_SECTION_FILL = PatternFill("solid", fgColor="E2E8F0")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _header(ws: Worksheet, columns: list[str], freeze: bool = True) -> None:
    ws.append(columns)
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    if freeze:  # a coordinate string: ws.cell() would create an empty row
        ws.freeze_panes = f"A{ws.max_row + 1}"


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _grid_sheet(ws: Worksheet, rows: list[dict], freeze: bool = True) -> None:
    """5-column grid with a shaded row per section, in registry order."""
    _header(ws, QUOTE_COLUMNS, freeze)
    section = None
    for r in rows:
        if r["section"] != section:
            section = r["section"]
            ws.append([SECTION_LABELS.get(section, section)])
            for cell in ws[ws.max_row]:
                cell.fill, cell.font = _SECTION_FILL, _BOLD
        ws.append([r["quote_field"], r["coverage_details"], r["proposed_details"],
                   r["source_status"], r["confidence"]])  # fmt: skip
        for cell in ws[ws.max_row]:
            cell.alignment = _WRAP
    _widths(ws, [42, 40, 40, 18, 12])


def build_quote_workbook(
    case: dict, grid: list[dict], sublimits: list[dict], notes: list[dict]
) -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Quote Summary"
    for label, value in [
        ("Company", case["company_name"]),
        ("Case Number", case["case_number"]),
        ("Type of Proposal", case["proposal_type"]),
        ("Rater Type", case["rater_type"]),
        ("Brokerage %", case["brokerage_pct"]),
        ("PIN Code", case["pincode"]),
        ("State", case["state"]),
        ("District", case["district"]),
        ("Zone", case["zone"]),
        ("Group Size", case["group_size"]),
        ("Number of Primary Members", case["primary_members_count"]),
        ("Flagging SME", "Yes" if case["sme_flag"] else "No"),
        ("Generated", datetime.now().strftime("%d-%m-%Y %H:%M")),
    ]:
        summary.append([label, value])
        summary.cell(row=summary.max_row, column=1).font = _BOLD
    summary.append([])
    _grid_sheet(summary, [r for r in grid if r["section"] not in TREATMENT_SECTIONS], freeze=False)

    _grid_sheet(
        wb.create_sheet("Hospitalization & Treatments"),
        [r for r in grid if r["section"] in TREATMENT_SECTIONS],
    )

    subs = wb.create_sheet("Procedure Sub-Limits")
    _header(subs, ["Procedure", "Option 1", "Option 2", "Option 3", "Option 4", "RFQ Amount",
                   "Matched Option", "Match Status", "Underwriter Approved",
                   "Approval Justification"])  # fmt: skip
    for s in sublimits:
        subs.append([
            s["procedure_name"], *s["options"].values(), s["extracted_amount_raw"],
            s["matched_option"], s["match_status"], "Yes" if s["is_underwriter_approved"] else "No",
            s["approval_justification"],
        ])  # fmt: skip
        for cell in subs[subs.max_row][1:5]:
            cell.number_format = "#,##,##0"
    _widths(subs, [48, 12, 12, 12, 12, 18, 16, 28, 12, 48])

    notes_ws = wb.create_sheet("Underwriter Notes")
    _header(notes_ws, ["Timestamp", "Reviewer", "Action", "Target", "Justification"])
    for n in notes:
        notes_ws.append([n["timestamp"].strftime("%d-%m-%Y %H:%M"), n["user_email"],
                         n["action"], n["target_entity"], n["justification_reason"]])  # fmt: skip
        notes_ws.cell(row=notes_ws.max_row, column=5).alignment = _WRAP
    _widths(notes_ws, [18, 32, 22, 48, 64])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
