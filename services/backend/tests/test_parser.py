import io

import docx
import pymupdf
import pytest

from app.services.parser import UnsupportedDocumentError, parse_document
from app.services.storage import PasswordRequiredException

from .llm_fakes import FIXTURES


def make_pdf(user_pw: str | None = None) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Room rent is capped at 1% of SI for normal room.")
    page.insert_htmlbox(
        pymupdf.Rect(72, 120, 520, 260),
        "<table style='border-collapse:collapse' border='1'>"
        "<tr><td style='border:1px solid black'>Procedure</td>"
        "<td style='border:1px solid black'>Limit</td></tr>"
        "<tr><td style='border:1px solid black'>CABG</td>"
        "<td style='border:1px solid black'>Rs. 1,20,000</td></tr></table>",
    )
    doc.new_page().insert_text((72, 72), "Second page: maternity covered.")
    if user_pw:
        return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=user_pw, owner_pw="o")
    return doc.tobytes()


def test_pdf_has_page_markers_text_and_tables():
    parsed = parse_document(make_pdf(), "rfq.pdf")
    md = parsed.markdown
    assert (parsed.pages, parsed.engine) == (2, "pymupdf")
    assert "[[File: rfq.pdf | Page 1]]" in md and "[[File: rfq.pdf | Page 2]]" in md
    assert "1% of SI for normal room" in md
    assert "|" in md and "CABG" in md and "1,20,000" in md
    assert md.index("Page 2]]") < md.index("maternity covered")


def test_locked_pdf_requires_password_and_decrypts_in_memory():
    data = make_pdf("s3cret")
    with pytest.raises(PasswordRequiredException):
        parse_document(data, "rfq.pdf")
    assert "maternity covered" in parse_document(data, "rfq.pdf", "s3cret").markdown


def test_docx_via_docling():
    d = docx.Document()
    d.add_heading("Coverage Requirements", level=1)
    d.add_paragraph("Pre-existing diseases covered from Day 1.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Procedure", "Limit"
    table.cell(1, 0).text, table.cell(1, 1).text = "CABG", "Rs. 1,20,000"
    buf = io.BytesIO()
    d.save(buf)
    parsed = parse_document(buf.getvalue(), "rfq.docx")
    assert parsed.engine == "docling"
    assert "# Coverage Requirements" in parsed.markdown
    assert "| CABG" in parsed.markdown


def test_text_and_markdown():
    parsed = parse_document((FIXTURES / "rfq_apex.md").read_bytes(), "rfq_apex.md")
    assert parsed.markdown.startswith("[[File: rfq_apex.md | Page 1]]")
    assert "ICICI Lombard" in parsed.markdown


def test_unsupported_type():
    with pytest.raises(UnsupportedDocumentError):
        parse_document(b"\x00\x01", "photo.png")
