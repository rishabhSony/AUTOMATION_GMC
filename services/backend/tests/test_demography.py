import io
import time
from datetime import date

import openpyxl
import polars as pl
import pytest

from app.services import demography as d

BASIS = date(2026, 10, 29)


def census(rows: list[list], headers: list[str] | None = None) -> pl.DataFrame:
    headers = headers or ["Emp ID", "Member Name", "Relation", "D.O.B.", "Age", "Sex", "SI"]
    return pl.DataFrame(
        [[None if v is None else str(v) for v in r] for r in rows],
        schema=headers,
        orient="row",
    )


def run(rows: list[list], headers: list[str] | None = None, basis: date = BASIS):
    return d.process_census(census(rows, headers), basis)


def exception_types(result: d.DemographyResult, row: int) -> set[str]:
    return set(
        result.exceptions.filter(pl.col("source_row_number") == row)["exception_type"].to_list()
    )


# ---------------------------------------------------------------------------
# Header aliases & normalization
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("header", "canonical"),
    [
        ("Emp ID", "employee_code"),
        ("ECODE", "employee_code"),
        ("Staff_Id", "employee_code"),
        ("Insured Name", "member_name"),
        ("Dep Relationship", "relationship"),
        ("D.O.B.", "dob"),
        ("DOB (DD/MM/YYYY)", "dob"),
        ("Completed Age", "age"),
        ("M/F", "gender"),
        ("sum_insured", "sum_insured"),
        ("Base SI", "sum_insured"),
    ],
)
def test_header_aliases(header, canonical):
    assert d.resolve_headers([header]) == {canonical: header}


def test_member_name_preferred_over_employee_name():
    mapping = d.resolve_headers(["Employee Name", "Member Name", "Relation", "DOB"])
    assert mapping["member_name"] == "Member Name"


def test_census_without_relationship_is_rejected():
    with pytest.raises(d.CensusFormatError):
        run([["E1", "A", "1990-01-01"]], headers=["Emp ID", "Name", "DOB"])


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Self", "employee"),
        ("WIFE", "spouse"),
        ("Husband", "spouse"),
        ("daughter", "child"),
        ("Dependent Child", "child"),
        ("Mother", "parent"),
        ("Father-in-law", "parent_in_law"),
        ("MIL", "parent_in_law"),
    ],
)
def test_relationship_normalization(raw, expected):
    result = run([["E1", "A", raw, "01-01-1990", None, "M", None]])
    assert result.records["relationship"].to_list() == [expected]


@pytest.mark.parametrize(
    "dob",
    [
        "14/06/1992",
        "14-06-1992",
        "14.06.1992",
        "1992-06-14",
        "1992-06-14 00:00:00",
        "14-Jun-1992",
        "33769",
        "06/30/1992",
    ],
)
def test_date_formats(dob):
    result = run([["E1", "A", "Self", dob, None, "M", None]])
    parsed = result.records["dob"][0]
    expected = date(1992, 6, 30) if dob == "06/30/1992" else date(1992, 6, 14)
    assert parsed == expected


def test_sum_insured_cleaning():
    result = run([["E1", "A", "Self", "01-01-1990", None, "M", "₹ 5,00,000/-"]])
    assert result.records["sum_insured"][0] == 500000.0


# ---------------------------------------------------------------------------
# Ages
# ---------------------------------------------------------------------------


def test_age_formula_uses_365_25_floor():
    # 18 calendar years minus one day → 17 by ⌊days / 365.25⌋
    result = run([["E1", "A", "Self", "29-02-2000", None, "M", None]], basis=date(2018, 2, 28))
    assert result.records["calculated_age"][0] == 17


def test_basis_toggle_changes_age():
    rows = [["E1", "A", "Self", "15-11-1990", None, "M", None]]
    assert run(rows, basis=date(2026, 11, 1)).records["calculated_age"][0] == 35
    assert run(rows, basis=date(2026, 11, 30)).records["calculated_age"][0] == 36


def test_source_age_fallback_when_dob_missing():
    result = run([["E1", "A", "Self", None, "41", "M", None]])
    assert result.records["calculated_age"][0] == 41


def test_missing_age_and_dob_is_critical_and_not_stored():
    result = run([["E1", "A", "Self", None, None, "M", None]])
    assert result.records.height == 0
    exc = result.exceptions.row(0, named=True)
    assert exc["exception_type"] == "missing_age_and_dob"
    assert exc["severity"] == "critical"
    assert "CRITICAL_ERROR: Missing Age and DOB" in exc["error_message"]


def test_age_dob_mismatch_over_one_year():
    result = run(
        [
            ["E1", "A", "Self", "01-01-1990", "36", "M", None],  # calculated 36 → ok
            ["E2", "B", "Self", "01-01-1990", "38", "M", None],  # off by 2 → flagged
        ]
    )
    assert "age_dob_mismatch" not in exception_types(result, 2)
    assert "age_dob_mismatch" in exception_types(result, 3)


def test_invalid_dob_falls_back_to_source_age():
    result = run([["E1", "A", "Self", "31-31-1990", "30", "M", None]])
    assert result.records["calculated_age"][0] == 30
    assert "invalid_dob" in exception_types(result, 2)


# ---------------------------------------------------------------------------
# Integrity rules & metrics
# ---------------------------------------------------------------------------


def test_integrity_flags():
    result = run(
        [
            ["E1", "Teen Worker", "Self", "01-01-2010", None, "M", None],  # row 2: under 18
            ["E9", "Lone Spouse", "Spouse", "01-01-1990", None, "F", None],  # row 3: orphan
            ["E1", "Dup Kid", "Son", "01-01-2015", None, "M", None],  # row 4
            ["E1", "Dup Kid", "Son", "01-01-2015", None, "M", None],  # row 5: duplicate of 4
            ["E1", None, "Daughter", "01-01-2016", None, None, None],  # row 6: missing fields
            ["E1", "X", "Cousin", "01-01-2000", None, "M", None],  # row 7: unmapped relation
        ]
    )
    assert "under18_employee" in exception_types(result, 2)
    assert exception_types(result, 3) == {"orphan_dependent"}
    assert "duplicate_member" in exception_types(result, 4)
    assert "duplicate_member" in exception_types(result, 5)
    assert exception_types(result, 6) == {"missing_required_field"}
    assert exception_types(result, 7) == {"unmapped_relationship"}

    stored = result.records.sort("source_row_number")
    assert stored["source_row_number"].to_list() == [2, 3, 4, 5, 6]
    assert stored["has_validation_exception"].all()
    assert stored.filter(pl.col("source_row_number") == 6)["member_name"][0] == ""


@pytest.mark.parametrize(("primaries", "sme"), [(250, True), (251, False)])
def test_group_metrics_and_sme_threshold(primaries, sme):
    rows = [[f"E{i}", f"N{i}", "Self", "01-01-1990", None, "M", None] for i in range(primaries)]
    rows.append(["E0", "Spouse", "Wife", "01-01-1991", None, "F", None])
    m = run(rows).metrics
    assert (m.group_size, m.primary_members_count, m.sme_flag) == (primaries + 1, primaries, sme)


def test_seven_column_export_order_and_format():
    result = run(
        [
            ["E1", "Amit Verma", "Self", "14/06/1992", None, "Male", "500000"],
            ["E1", "Pooja Verma", "Wife", "22/11/1995", None, "F", "500000"],
            ["E1", "Old", "Father-in-law", "01/01/1960", None, "M", None],
        ]
    )
    seven = d.to_seven_column(result.records)
    assert seven.columns == d.SEVEN_COLUMN_ORDER
    assert seven.row(0) == (34, "14-06-1992", 500000.0, "Employee", "Amit Verma", "E1", "Male")
    assert seven["Relationship"][2] == "Parent-in-law"

    wb = openpyxl.load_workbook(io.BytesIO(d.seven_column_excel_bytes(result.records)))
    header = [c.value for c in next(wb.active.iter_rows(max_row=1))]
    assert header == d.SEVEN_COLUMN_ORDER


def test_reads_xlsx_with_native_dates():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Employee Code", "Name", "Relationship", "Date of Birth", "Gender", "Sum Insured"])
    ws.append(["E1", "Amit", "Self", date(1992, 6, 14), "M", 500000])
    buf = io.BytesIO()
    wb.save(buf)
    raw = d.read_census(io.BytesIO(buf.getvalue()), "census.xlsx")
    result = d.process_census(raw, BASIS)
    assert result.records["dob"][0] == date(1992, 6, 14)
    assert result.records["calculated_age"][0] == 34


# ---------------------------------------------------------------------------
# Performance: 50,000-row synthetic census (Plan Phase N+1)
# ---------------------------------------------------------------------------


def synthetic_census_csv(n: int) -> bytes:
    """Families of 4 (Self, Wife, Son, Mother) with mixed date formats and some source ages."""
    i = pl.int_range(0, n)
    day = (i % 28 + 1).cast(pl.String).str.zfill(2)
    month = (i % 12 + 1).cast(pl.String).str.zfill(2)
    year = (1950 + i % 60).cast(pl.String)
    frame = pl.select(
        pl.format("EMP{}", (i // 4).cast(pl.String).str.zfill(6)).alias("Emp ID"),
        pl.format("Member {}", i).alias("Insured Name"),
        pl.lit(pl.Series(["Self", "Wife", "Son", "Mother"])).gather(i % 4).alias("Relation"),
        pl.when(i % 3 == 0)
        .then(pl.format("{}/{}/{}", day, month, year))
        .otherwise(pl.format("{}-{}-{}", year, month, day))
        .alias("DOB"),
        pl.when(i % 10 == 0).then((2026 - 1950 - i % 60).cast(pl.String)).alias("Age"),
        pl.lit(pl.Series(["M", "F", "M", "F"])).gather(i % 4).alias("Gender"),
        pl.lit("5,00,000").alias("Sum Insured"),
    )
    return frame.write_csv().encode()


def test_50k_rows_under_two_seconds():
    data = synthetic_census_csv(50_000)
    start = time.perf_counter()
    raw = d.read_census(io.BytesIO(data), "census.csv")
    result = d.process_census(raw, BASIS)
    elapsed = time.perf_counter() - start

    assert result.metrics.group_size == 50_000
    assert result.metrics.primary_members_count == 12_500
    assert result.metrics.sme_flag is False
    assert result.records["calculated_age"].null_count() == 0
    assert elapsed < 2.0, f"50k-row census took {elapsed:.2f}s"
